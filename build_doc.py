from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path
import os, textwrap

root = Path(r'C:\Users\44244\Desktop\guanlihoutai')
out = root / 'docs' / 'word_assets'
out.mkdir(parents=True, exist_ok=True)
font_path = r'C:\Windows\Fonts\msyh.ttc'
font_bold_path = r'C:\Windows\Fonts\msyh.ttc'

def font(size, bold=False):
    return ImageFont.truetype(font_bold_path if bold else font_path, size)

def rounded(draw, xy, fill, outline='#CBD5E1', radius=18, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)

def arrow(draw, a, b, color='#64748B', width=4):
    draw.line([a,b], fill=color, width=width)
    import math
    ang=math.atan2(b[1]-a[1], b[0]-a[0])
    L=16
    p1=(b[0]-L*math.cos(ang-0.5), b[1]-L*math.sin(ang-0.5))
    p2=(b[0]-L*math.cos(ang+0.5), b[1]-L*math.sin(ang+0.5))
    draw.polygon([b,p1,p2], fill=color)

def centered(draw, box, text, f, fill='#0F172A', spacing=6):
    x1,y1,x2,y2=box
    lines=[]
    for para in text.split('\n'):
        lines.extend(textwrap.wrap(para, width=max(8,int((x2-x1)/(f.size*0.95)))))
    heights=[draw.textbbox((0,0),ln,font=f)[3] for ln in lines]
    total=sum(heights)+spacing*(len(lines)-1)
    y=y1+(y2-y1-total)/2
    for ln,h in zip(lines,heights):
        bb=draw.textbbox((0,0),ln,font=f); w=bb[2]-bb[0]
        draw.text((x1+(x2-x1-w)/2,y),ln,font=f,fill=fill)
        y+=h+spacing

def make_overall(path):
    W,H=2200,1550
    im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im)
    title=font(46,True); sub=font(28); boxf=font(27,True); small=font(24)
    d.text((70,45),'星河游戏俱乐部完整业务主链路',font=title,fill='#0F172A')
    d.text((72,105),'实线表示当前已打通；虚线表示待接入或需要补齐的产品能力',font=sub,fill='#475569')
    nodes=[
      ('用户进入\n公开首页 / 注册登录','#E0F2FE','#0369A1'),
      ('个人中心\n余额 / 点单 / 售后','#E0F2FE','#0369A1'),
      ('客服确认需求\n游戏 商品 时长 支付','#DCFCE7','#15803D'),
      ('创建订单\n服务端计算金额','#DCFCE7','#15803D'),
      ('派单或接单\n游戏 等级 在线状态','#DCFCE7','#15803D'),
      ('确认接单\n开始服务','#DCFCE7','#15803D'),
      ('提交完单说明\n客服或售后验收','#DCFCE7','#15803D'),
      ('收益入账\n抽佣比例按派单锁定','#DCFCE7','#15803D'),
      ('申请提现\n冻结可提现余额','#FEF3C7','#B45309'),
      ('财务审核\n待线下打款','#FEF3C7','#B45309'),
      ('登记打款流水\n已打款 / 审计留痕','#FEF3C7','#B45309'),
    ]
    x0,y0,bw,bh,gap=70,220,350,150,45
    coords=[]
    for i,(txt,fill,edge) in enumerate(nodes):
        row=i//4; col=i%4
        x=x0+col*(bw+gap); y=y0+row*(bh+90)
        if i>=8: y += 30
        rounded(d,(x,y,x+bw,y+bh),fill,edge,22,4); centered(d,(x+18,y+12,x+bw-18,y+bh-12),txt,boxf)
        coords.append((x,y,x+bw,y+bh))
    for i in range(len(coords)-1):
        a=coords[i]; b=coords[i+1]
        if i%4 != 3:
            arrow(d,(a[2]+5,(a[1]+a[3])//2),(b[0]-5,(b[1]+b[3])//2))
        else:
            arrow(d,((a[0]+a[2])//2,a[3]+5),((b[0]+b[2])//2,b[1]-5))
    # side dashed gaps
    gaps=[('待接入：真实用户中心 / 微信 / 游戏平台',70,1130,650,1280),('待接入：实时消息、未读推送、SLA 自动催办',770,1130,1350,1280),('待接入：支付回调、退款回调、银行打款接口',1470,1130,2130,1280)]
    for txt,x1,y1,x2,y2 in gaps:
        d.rounded_rectangle((x1,y1,x2,y2),radius=18,fill='#F8FAFC',outline='#94A3B8',width=4)
        for xx in range(x1,x2,24): d.line((xx,y1, min(xx+12,x2),y1),fill='#94A3B8',width=4)
        centered(d,(x1+15,y1+10,x2-15,y2-10),txt,small,fill='#475569')
    im.save(path,dpi=(220,220))

def make_roles(path):
    W,H=2200,1450
    im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im)
    title=font(44,True); sub=font(27); head=font(30,True); body=font(25)
    d.text((70,45),'角色与后台工作区链路',font=title,fill='#0F172A')
    d.text((72,105),'角色由服务端校验；角色变化会撤销旧会话。陪玩身份与职责分开管理。',font=sub,fill='#475569')
    rounded(d,(70,220,560,430),'#E0F2FE','#0369A1',22,4); centered(d,(100,250,530,400),'用户池\n已有账号与个人资料',head)
    roles=[('最高负责人','首页、成员、财务、订单、审计','#DBEAFE','#1D4ED8'),('客服','会话、创建订单、派单、完单验收','#DCFCE7','#15803D'),('财务','充值审核、流水、提现结算','#FEF3C7','#B45309'),('考官 / 售后','考核质检、退款、订单验收','#FCE7F3','#BE185D'),('陪玩','接单大厅、本人订单、收益提现','#EDE9FE','#7C3AED'),('普通成员','个人资料；个人订单/钱包入口需补齐权限链路','#F1F5F9','#64748B')]
    x,y=760,190; bw,bh=540,145
    for i,(name,txt,fill,edge) in enumerate(roles):
        yy=y+i*190
        rounded(d,(x,yy,x+bw,yy+bh),fill,edge,20,4)
        d.text((x+25,yy+20),name,font=head,fill='#0F172A')
        centered(d,(x+25,yy+60,x+bw-25,yy+bh-18),txt,body,fill='#334155')
        arrow(d,(560,325),(x-10,yy+bh//2),color='#64748B')
    rounded(d,(1450,220,2120,470),'#F8FAFC','#94A3B8',22,4)
    centered(d,(1485,250,2085,440),'陪玩准入\n押金 + 等级 + 游戏\n冻结 / 解冻 / 订单冲突校验',head,fill='#0F172A')
    rounded(d,(1450,560,2120,810),'#F8FAFC','#94A3B8',22,4)
    centered(d,(1485,590,2085,780),'订单约束\n版本检查、重复派单拦截\n同一成员不可并行开始两单',head,fill='#0F172A')
    rounded(d,(1450,900,2120,1150),'#FFF7ED','#C2410C',22,4)
    centered(d,(1485,930,2085,1120),'待修复断链\n普通成员 pages 只有 overview\n前端会话未调用消息发送接口',head,fill='#9A3412')
    arrow(d,(1330,270),(1450,345)); arrow(d,(1330,460),(1450,685)); arrow(d,(1330,840),(1450,1025))
    im.save(path,dpi=(220,220))

def make_roadmap(path):
    W,H=2200,1350
    im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im)
    title=font(44,True); sub=font(27); hf=font(31,True); bf=font(25)
    d.text((70,45),'未完成项与生产落地顺序',font=title,fill='#0F172A')
    d.text((72,105),'按“先补断链、再接外部系统、最后做生产治理”的顺序推进。',font=sub,fill='#475569')
    cols=[('第一阶段\n补齐产品断链','#FFF7ED','#C2410C',['普通成员个人页：继承用户能力或移除死入口','会话增加消息输入、发送、未读状态','游戏与商品配置增加增删改、上下架、价格维护','财务列表改为真实资金明细来源']),('第二阶段\n接入外部系统','#EFF6FF','#1D4ED8',['真实用户中心 / 微信 / 游戏平台同步','充值支付回调、签名校验、自动对账','第三方原路退款与退款回调','提现银行或支付接口']),('第三阶段\n生产治理与运营','#F0FDF4','#15803D',['SLA 自动催办、消息通知、异常告警','定时报表与完整明细导出','HTTPS、二次验证、密码策略','备份恢复演练、限流、监控、敏感字段脱敏'])]
    x=70; y=220; cw=650; gap=55
    for i,(name,fill,edge,items) in enumerate(cols):
        xx=x+i*(cw+gap)
        rounded(d,(xx,y,xx+cw,y+850),fill,edge,24,4)
        centered(d,(xx+25,y+25,xx+cw-25,y+110),name,hf,fill='#0F172A')
        yy=y+155
        for item in items:
            rounded(d,(xx+35,yy,xx+cw-35,yy+130),'white',edge,16,3)
            centered(d,(xx+55,yy+18,xx+cw-55,yy+112),item,bf,fill='#334155')
            yy+=165
        if i<2: arrow(d,(xx+cw+10,y+425),(xx+cw+gap-10,y+425),color='#64748B',width=5)
    rounded(d,(70,1120,2125,1245),'#F8FAFC','#94A3B8',18,3)
    centered(d,(100,1140,2095,1225),'当前验证基线：node --test tests/*.test.mjs 共 34 项通过；说明后端主状态机已具备继续补齐前端与外部集成的基础。',bf,fill='#334155')
    im.save(path,dpi=(220,220))

make_overall(out/'overall.png'); make_roles(out/'roles.png'); make_roadmap(out/'roadmap.png')

# doc helpers

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9', sz='8'):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); borders=tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge; el=borders.find(qn(tag))
        if el is None: el=OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),sz); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)

def set_run_font(run, name='Microsoft YaHei', size=10.5, bold=False, color='000000'):
    run.font.name=name; run._element.rPr.rFonts.set(qn('w:ascii'),name); run._element.rPr.rFonts.set(qn('w:hAnsi'),name); run._element.rPr.rFonts.set(qn('w:eastAsia'),name); run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor.from_string(color)

def add_p(doc, text='', style=None, bold_prefix=None):
    p=doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r=p.add_run(bold_prefix); set_run_font(r,bold=True); r2=p.add_run(text[len(bold_prefix):]); set_run_font(r2)
    else:
        r=p.add_run(text); set_run_font(r)
    return p

def add_bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet'); r=p.add_run(item); set_run_font(r)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
styles=doc.styles
for s in ['Normal','Title','Heading 1','Heading 2','List Bullet']:
    st=styles[s]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); st.font.color.rgb=RGBColor(0,0,0)
styles['Normal'].font.size=Pt(10.5); styles['Normal'].paragraph_format.space_after=Pt(6)
styles['Title'].font.size=Pt(24); styles['Title'].font.bold=True
styles['Heading 1'].font.size=Pt(16); styles['Heading 1'].font.bold=True; styles['Heading 1'].paragraph_format.space_before=Pt(10); styles['Heading 1'].paragraph_format.space_after=Pt(6)
styles['Heading 2'].font.size=Pt(12.5); styles['Heading 2'].font.bold=True

p=doc.add_paragraph(style='Title'); r=p.add_run('星河游戏俱乐部网站全链路整理'); set_run_font(r,size=24,bold=True)
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('已完成能力、当前断点与生产落地路线'); set_run_font(r,size=12,color='475569')
add_p(doc,'本文件把当前网站从用户进入、订单履约、资金结算到后台治理的完整链路统一整理，并区分已经打通的功能与仍需补齐的能力。当前后端主状态机、权限校验和资金状态流转已经具备继续完善的基础；需要优先处理的是普通成员个人入口、会话消息发送、配置维护页面与真实外部系统接入。')

doc.add_heading('一 完整业务主链路', level=1)
doc.add_picture(str(out/'overall.png'), width=Inches(6.8)); p=doc.add_paragraph('图 1  网站从用户进入到订单履约和提现结算的主链路'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; [set_run_font(r,size=9,color='64748B') for r in p.runs]
add_p(doc,'实线框表示当前已经在代码和测试中形成闭环的步骤。虚线框表示当前范围内明确存在、但还没有接入真实系统或自动化能力的部分。')

doc.add_heading('二 角色与后台工作区', level=1)
doc.add_picture(str(out/'roles.png'), width=Inches(6.8)); p=doc.add_paragraph('图 2  角色、陪玩准入和后台权限关系'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; [set_run_font(r,size=9,color='64748B') for r in p.runs]
add_p(doc,'最高负责人负责配置和治理；客服、财务、考官、售后和陪玩分别承担业务环节。陪玩身份与职责分开管理，设为陪玩时需要押金、等级和至少一个游戏。角色由服务端逐项限制，角色变化会撤销旧会话。')

doc.add_heading('三 当前已经完成', level=1)
add_bullets(doc,[
'用户先进入俱乐部用户池，再由负责人使用已有用户 ID 加入成员，避免重复创建登录账号。',
'角色权限、订单归属、会话有效性和管理接口由服务端校验；停用或改职责会撤销旧会话。',
'陪玩准入包含押金、等级、游戏和冻结状态；等级决定接单门槛，游戏抽佣比例在派单时锁定。',
'订单支持版本检查、重复派单拦截、服务冲突拦截、在线状态和游戏匹配校验。',
'订单金额由服务端按商品和时长计算，余额支付即时扣减；验收通过后才入账。',
'充值审核、退款申请、部分退款额度、提现冻结、驳回返还、线下打款登记和审计记录已经形成状态机。',
'经营汇总、趋势、陪玩榜、用户榜、订单榜支持日期、游戏、Tag、金额或笔数筛选，并支持 CSV 汇总导出。',
'当前测试基线为 node --test tests/*.test.mjs，共 34 项通过。'
])

doc.add_heading('四 已发现的断链与风险', level=1)
rows=[
('普通成员个人入口','前端导航列出我的点单、售后记录、我的钱包，但服务端 member 角色 pages 只有 overview，进入后会被回退。','高','让普通成员继承用户个人能力并补齐 workspace 数据；或删除这些入口并明确限制。'),
('会话消息发送','服务端存在消息接口，但会话弹窗没有消息输入和发送调用，目前只能查看历史消息、保存内部跟进。','高','增加消息输入、发送、未读状态和客服/用户双方消息展示。'),
('游戏与商品配置','catalog 页面当前以展示为主，缺少增删改、上下架和价格维护链路。','中','补齐配置 API、表单、校验、变更审计与订单引用保护。'),
('财务列表数据来源','部分财务列表仍使用历史示例或前端推算，真正资金明细在资金流水页。','中','统一以服务端流水和结算记录作为唯一数据源。'),
('外部支付与消息','用户中心、充值支付回调、原路退款、银行打款、实时消息和自动催办尚未接入。','高','按下一节路线分阶段接入，并为每个外部回调增加幂等号和对账。'),
('生产安全治理','生产 HTTPS、管理员二次验证、备份恢复、限流监控和敏感字段脱敏尚未完成。','高','上线前完成安全基线、备份演练和关键操作审计。')]
t=doc.add_table(rows=1, cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
for i,h in enumerate(['事项','当前表现','优先级','完善建议']):
    c=t.rows[0].cells[i]; c.text=''; set_cell_shading(c,'1F4E78'); set_cell_border(c); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(h); set_run_font(r,size=10,bold=True,color='FFFFFF'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
for rowdata in rows:
    cells=t.add_row().cells
    for i,val in enumerate(rowdata):
        cells[i].text=''; set_cell_border(cells[i]); set_cell_shading(cells[i], 'F7FAFC' if len(t.rows)%2==0 else 'FFFFFF'); p=cells[i].paragraphs[0]; r=p.add_run(val); set_run_font(r,size=9.5); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER


doc.add_heading('五 建议的生产落地顺序', level=1)
doc.add_picture(str(out/'roadmap.png'), width=Inches(6.8)); p=doc.add_paragraph('图 3  未完成项按阶段推进的建议路线'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; [set_run_font(r,size=9,color='64748B') for r in p.runs]
add_bullets(doc,[
'第一阶段先补齐前端和权限断链，确保普通成员、消息会话、游戏商品配置和财务明细都能从入口走到结果。',
'第二阶段接入用户中心、支付、退款和提现等外部系统，所有回调使用幂等号、签名校验和自动对账。',
'第三阶段建设生产治理，包括通知、SLA 自动催办、异常告警、定时报表、HTTPS、二次验证、备份恢复和监控。'
])

doc.add_heading('六 交付判断', level=1)
add_p(doc,'当前版本可以作为本地演示和内部流程评审版本：核心订单、角色、资金和统计链路已经可验证；要进入真实运营，还需要完成断链修复、外部支付与消息接入，以及生产安全和数据治理。')

# footer
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=footer.add_run('星河游戏俱乐部后台链路整理'); set_run_font(r,size=9,color='64748B')

path=root/'docs'/'星河游戏俱乐部全链路整理.docx'; doc.save(path)
print(path)
